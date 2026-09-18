from __future__ import annotations

from pathlib import Path
from typing import Optional

from ..providers import AnalysisResult


def render_redline_markdown(result: AnalysisResult) -> str:
    """A focused, negotiation-ready view: only the clauses that need changes."""
    flagged = [c for c in result.clauses if c.get("is_red_flag") or c.get("negotiation_suggestion")]
    lines = [f"# Redline Suggestions - {result.contract_type}", ""]
    if not flagged:
        lines.append("No clauses were flagged for renegotiation.")
        return "\n".join(lines)

    for c in flagged:
        lines += [
            f"## {c.get('name')}  (risk {c.get('score')}/10)",
            "",
            f"**Issue:** {c.get('finding', '')}",
            "",
            f"**Suggested replacement language:**",
            f"> {c.get('negotiation_suggestion') or 'Consult an attorney for specific replacement language.'}",
            "",
        ]
    return "\n".join(lines)


def write_redline_docx(result: AnalysisResult, out_path: str | Path) -> Optional[Path]:
    """
    Write a Word document with one heading + issue + suggestion per
    flagged clause, styled so it reads like a redline memo you can attach
    to an email back to the counter-party. Returns None (and does not
    raise) if python-docx is not installed, so this stays a purely
    optional feature.
    """
    try:
        import docx
        from docx.shared import Pt, RGBColor
    except ImportError:
        return None

    out_path = Path(out_path)
    document = docx.Document()
    document.add_heading(f"Redline Suggestions - {result.contract_type}", level=1)
    document.add_paragraph(f"Overall risk: {result.overall_risk}  |  Verdict: {result.verdict}")

    flagged = [c for c in result.clauses if c.get("is_red_flag") or c.get("negotiation_suggestion")]
    if not flagged:
        document.add_paragraph("No clauses were flagged for renegotiation.")
    for c in flagged:
        document.add_heading(f"{c.get('name')} (risk {c.get('score')}/10)", level=2)
        p = document.add_paragraph()
        run = p.add_run("Issue: ")
        run.bold = True
        p.add_run(c.get("finding", ""))

        p2 = document.add_paragraph()
        run2 = p2.add_run("Suggested replacement language: ")
        run2.bold = True
        run3 = p2.add_run(c.get("negotiation_suggestion") or "Consult an attorney for specific language.")
        run3.italic = True
        run3.font.color.rgb = RGBColor(0x1A, 0x73, 0x2E)

    document.add_paragraph()
    footer = document.add_paragraph(
        "Hermes Legal Advisor provides contract analysis, not legal advice. "
        "Always consult a qualified attorney before signing any contract."
    )
    footer.runs[0].font.size = Pt(8)

    document.save(str(out_path))
    return out_path
