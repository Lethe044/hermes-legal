"""
Document readers for Hermes Legal Advisor.

Real contracts rarely arrive as .txt files. This module adds first-class
support for PDF and DOCX so the tool is actually usable on documents
people receive by email, in addition to the plain text files used in the
sample/demo set.
"""

from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class UnsupportedFileError(ValueError):
    pass


def read_document(path: str | Path, allow_ocr: bool = True) -> str:
    """Read a contract file and return its plain text content.

    If a PDF has no extractable text (a scanned image PDF) and allow_ocr
    is True, automatically falls back to OCR via pytesseract + pdf2image
    if both are installed. This keeps plain text extraction fast as the
    default path and only pays the OCR cost when it's actually needed.
    """
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"File not found: {p}")

    suffix = p.suffix.lower()
    if suffix in (".txt", ".md"):
        return _read_txt(p)
    if suffix == ".pdf":
        return _read_pdf(p, allow_ocr=allow_ocr)
    if suffix == ".docx":
        return _read_docx(p)

    raise UnsupportedFileError(
        f"Unsupported file type '{suffix}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
    )


def _read_txt(p: Path) -> str:
    return p.read_text(encoding="utf-8", errors="replace")


def _read_pdf(p: Path, allow_ocr: bool = True) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ImportError(
            "Reading PDF files requires pypdf. Install with: pip install pypdf"
        ) from exc

    reader = PdfReader(str(p))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)
    combined = "\n\n".join(pages).strip()
    if combined:
        return combined

    if not allow_ocr:
        raise ValueError(
            f"No extractable text found in {p.name}. It may be a scanned image PDF "
            "that needs OCR before it can be analyzed."
        )

    ocr_text = _ocr_pdf(p)
    if not ocr_text.strip():
        raise ValueError(
            f"No extractable text found in {p.name}, and OCR also produced nothing usable. "
            "The scan quality may be too low, or the file may not contain readable text."
        )
    return ocr_text


def _ocr_pdf(p: Path) -> str:
    """OCR a scanned PDF page by page. Requires pytesseract, pdf2image, and
    the system tesseract and poppler binaries. Raises a clear, actionable
    error if any of those are missing rather than failing silently."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError as exc:
        raise ImportError(
            f"'{p.name}' appears to be a scanned PDF with no extractable text. "
            "OCR fallback requires: pip install pytesseract pdf2image, plus the "
            "system packages tesseract-ocr and poppler-utils."
        ) from exc

    try:
        images = convert_from_path(str(p))
    except Exception as exc:
        raise RuntimeError(
            f"Could not rasterize '{p.name}' for OCR. Is poppler-utils installed "
            "(the 'pdftoppm' command)?"
        ) from exc

    texts = []
    for image in images:
        texts.append(pytesseract.image_to_string(image))
    return "\n\n".join(texts).strip()


def _read_docx(p: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ImportError(
            "Reading DOCX files requires python-docx. Install with: pip install python-docx"
        ) from exc

    document = docx.Document(str(p))
    parts = [para.text for para in document.paragraphs if para.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)
